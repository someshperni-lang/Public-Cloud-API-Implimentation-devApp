#!/usr/bin/env python3
"""
Bulk SUBCONTRACTING Purchase Order upload (ME21N, item category L)
- API_PURCHASEORDER_PROCESS_SRV, SAP DEV.

Subcontracting only. Standard / account-assigned POs use
bulk_purchase_order_upload.py.

Item category
-------------
The API wants the INTERNAL key "3", not the external "L" shown in ME21N
(sending "L" fails with "Use a supported purchase order item category").
Blank, "L" or "3" in the sheet are all sent as "3"; anything else is rejected.

Two input layouts are accepted
------------------------------
1. Excel workbook with sheets Header / Items / ScheduleLines / Components
   (same layout as Subcontracting_PO_Upload.xlsx), linked by
   PO_Ref -> PurchaseOrderItem -> ScheduleLine.
2. One flat sheet (CSV or Excel): one row per component, with PO_Ref, item and
   schedule-line columns repeated. Header / item / schedule-line values are
   read from the first row of each PO / item / schedule line.

Both are turned into the same deep insert:
    A_PurchaseOrder -> to_PurchaseOrderItem -> to_ScheduleLine
                    -> to_SubcontractingComponent

Components
----------
Every item needs at least one component, unless ALLOW_BOM_EXPLOSION is True
(then items without components are sent with only a schedule line and SAP
explodes the subcontracting BOM).

    python bulk_subcontracting_po_upload.py --template xlsx
    python bulk_subcontracting_po_upload.py -f Subcontracting_PO_Upload.xlsx --dry-run
    python bulk_subcontracting_po_upload.py -f Subcontracting_PO_Upload.xlsx
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sap_common import (EntityBuilder, clean, load_metadata, pad, read_table,  # noqa: E402
                        rows_of, run_cli, service_url, validate_payload)

TITLE = "Subcontracting purchase orders"
SLUG = "subcontracting_po"
SERVICE = "API_PURCHASEORDER_PROCESS_SRV"
ENTITY_SET = "A_PurchaseOrder"
HEADER_TYPE = "A_PurchaseOrderType"
ITEM_TYPE = "A_PurchaseOrderItemType"
SCHED_TYPE = "A_PurchaseOrderScheduleLineType"
COMP_TYPE = "A_POSubcontractingComponentType"
BASE_URL = service_url(SERVICE)          # DEV: my440263
DOC_FIELD = "PurchaseOrder"
YEAR_FIELD = None
TCODE = "ME21N"
ICON = "🏭"
COLOR = "#7C3AED"
GROUP_LABEL = "PO_Ref"
ITEMS_LABEL = "Items"
DESCRIPTION = ("Subcontracting purchase orders (item category L) with the components you "
               "provide to the subcontractor. Upload the 4-sheet workbook "
               "(Header / Items / ScheduleLines / Components) or one flat sheet.")

SUBCONTRACTING_ITEM_CATEGORY = "3"          # internal key for external "L"
ACCEPTED_CATEGORY_INPUTS = {"", "L", "3"}
ALLOW_BOM_EXPLOSION = False


def count_items(p: dict) -> int:
    return len(p.get("to_PurchaseOrderItem", {}).get("results", []))


HEADER_FIELDS = ["CompanyCode", "PurchaseOrderType", "Supplier", "PurchasingOrganization",
                 "PurchasingGroup", "DocumentCurrency", "PurchaseOrderDate", "PaymentTerms",
                 "IncotermsClassification", "IncotermsLocation1", "CorrespncExternalReference",
                 "Language"]
ITEM_FIELDS = ["PurchaseOrderItem", "Material", "Plant", "StorageLocation", "OrderQuantity",
               "PurchaseOrderQuantityUnit", "NetPriceAmount", "NetPriceQuantity", "TaxCode",
               "PurchaseOrderItemCategory", "GoodsReceiptIsExpected", "InvoiceIsExpected",
               "InvoiceIsGoodsReceiptBased", "PurchasingInfoRecord", "PurchaseOrderItemText"]
SCHED_MAP = {"ScheduleLine": "ScheduleLine", "ScheduleLineDeliveryDate": "DeliveryDate",
             "ScheduleLineOrderQuantity": "ScheduleLineQuantity"}
COMP_MAP = {"ReservationItem": "ComponentReservationItem", "Material": "ComponentMaterial",
            "RequirementDate": "ComponentRequirementDate",
            "QuantityInEntryUnit": "ComponentQuantity", "EntryUnit": "ComponentUnit",
            "Plant": "ComponentPlant"}

COLUMNS = (["PO_Ref"] + HEADER_FIELDS + ITEM_FIELDS + list(SCHED_MAP.values())
           + list(COMP_MAP.values()))

COLUMN_NOTES = {
    "PO_Ref": "Required. Your own key - rows with the same value form one PO.",
    "CompanyCode": "Required (first row of the PO).",
    "PurchaseOrderType": "Required, e.g. ZSUB.",
    "Supplier": "Required. The subcontractor.",
    "PurchasingOrganization": "Required.",
    "PurchasingGroup": "Required, e.g. 001.",
    "DocumentCurrency": "Required, e.g. INR.",
    "PurchaseOrderItem": "Item number, e.g. 10 (padded to 00010). Defaults to 10, 20, 30...",
    "Material": "Required. The finished good bought back from the subcontractor.",
    "Plant": "Required.",
    "OrderQuantity": "Required.",
    "PurchaseOrderQuantityUnit": "Required, e.g. PC.",
    "NetPriceAmount": "Subcontracting (processing) price.",
    "TaxCode": "Tax code, e.g. G4.",
    "PurchaseOrderItemCategory": "Leave blank, L or 3 - always sent as 3 (subcontracting).",
    "GoodsReceiptIsExpected": "TRUE / FALSE.",
    "InvoiceIsExpected": "TRUE / FALSE.",
    "ScheduleLine": "Defaults to 0001.",
    "DeliveryDate": "Required. Delivery date of the schedule line.",
    "ScheduleLineQuantity": "Defaults to OrderQuantity.",
    "ComponentReservationItem": "Optional. Numbered 0001, 0002 ... per schedule line when blank.",
    "ComponentMaterial": "Required (one row per component) unless BOM explosion is allowed.",
    "ComponentRequirementDate": "Defaults to DeliveryDate.",
    "ComponentQuantity": "Required with ComponentMaterial.",
    "ComponentUnit": "Required with ComponentMaterial, e.g. KG.",
    "ComponentPlant": "Defaults to the item Plant.",
}

SAMPLE_ROWS = [
    {"PO_Ref": "PO1", "CompanyCode": "1810", "PurchaseOrderType": "ZSUB", "Supplier": "18300010",
     "PurchasingOrganization": "1810", "PurchasingGroup": "001", "DocumentCurrency": "INR",
     "PurchaseOrderItem": "10", "Material": "FG228", "Plant": "1810", "OrderQuantity": "100",
     "PurchaseOrderQuantityUnit": "PC", "NetPriceAmount": "250.00", "NetPriceQuantity": "1",
     "TaxCode": "G4", "PurchaseOrderItemCategory": "3", "GoodsReceiptIsExpected": "TRUE",
     "InvoiceIsExpected": "TRUE", "ScheduleLine": "1", "DeliveryDate": "2026-10-20",
     "ScheduleLineQuantity": "100", "ComponentMaterial": "RM019", "ComponentQuantity": "100",
     "ComponentUnit": "KG", "ComponentPlant": "1810"},
]

# ----------------------------------------------------------------------------
# 4-sheet workbook <-> flat rows
# ----------------------------------------------------------------------------
SHEETS = {
    "Header": ["PO_Ref"] + HEADER_FIELDS,
    "Items": ["PO_Ref"] + ITEM_FIELDS,
    "ScheduleLines": ["PO_Ref", "PurchaseOrderItem", "ScheduleLine", "ScheduleLineDeliveryDate",
                      "ScheduleLineOrderQuantity"],
    "Components": ["PO_Ref", "PurchaseOrderItem", "ScheduleLine", "ReservationItem", "Material",
                   "RequirementDate", "QuantityInEntryUnit", "EntryUnit", "Plant"],
}
SCHED_SHEET_TO_FLAT = {"ScheduleLineDeliveryDate": "DeliveryDate",
                       "ScheduleLineOrderQuantity": "ScheduleLineQuantity"}
COMP_SHEET_TO_FLAT = {v_sap: v_flat for v_sap, v_flat in COMP_MAP.items()}


def _sheet(xl, name) -> list[dict]:
    if name not in xl.sheet_names:
        return []
    df = pd.read_excel(xl, sheet_name=name, dtype=str, keep_default_na=False)
    df.columns = [str(c).strip() for c in df.columns]
    rows = [{k: clean(v) for k, v in r.items()} for r in df.to_dict(orient="records")]
    return [r for r in rows if any(r.values())]


def workbook_to_flat(raw: bytes) -> pd.DataFrame:
    """Header/Items/ScheduleLines/Components -> one row per component."""
    xl = pd.ExcelFile(io.BytesIO(raw))
    headers = _sheet(xl, "Header")
    items = _sheet(xl, "Items")
    scheds = _sheet(xl, "ScheduleLines")
    comps = _sheet(xl, "Components")

    def key_item(r):
        return (r.get("PO_Ref", ""), pad("PurchaseOrderItem", r.get("PurchaseOrderItem", "")))

    def key_sched(r):
        return key_item(r) + (pad("ScheduleLine", r.get("ScheduleLine", "") or "1"),)

    out = []
    for h in headers:
        hrow = {c: h.get(c, "") for c in ["PO_Ref"] + HEADER_FIELDS}
        po_items = [i for i in items if i.get("PO_Ref") == h.get("PO_Ref")]
        if not po_items:
            out.append(hrow)
            continue
        for it in po_items:
            irow = {**hrow, **{c: it.get(c, "") for c in ITEM_FIELDS}}
            it_scheds = [s for s in scheds if key_item(s) == key_item(it)]
            if not it_scheds:
                out.append(irow)
                continue
            for s in it_scheds:
                srow = {**irow, "ScheduleLine": s.get("ScheduleLine", "")}
                srow.update({flat: s.get(col, "") for col, flat in SCHED_SHEET_TO_FLAT.items()})
                s_comps = [c for c in comps if key_sched(c) == key_sched(s)]
                if not s_comps:
                    out.append(srow)
                for c in s_comps:
                    out.append({**srow, **{flat: c.get(col, "")
                                           for col, flat in COMP_SHEET_TO_FLAT.items()}})
    # components / items pointing at a PO_Ref that has no header row
    known = {h.get("PO_Ref") for h in headers}
    for r in items + comps:
        if r.get("PO_Ref") and r["PO_Ref"] not in known:
            out.append({"PO_Ref": r["PO_Ref"]})
            known.add(r["PO_Ref"])
    return pd.DataFrame(out, columns=COLUMNS).fillna("")


def read_upload(source, filename: str) -> pd.DataFrame:
    """Accept the 4-sheet workbook or a flat CSV / Excel sheet."""
    if filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raw = source if isinstance(source, (bytes, bytearray)) else Path(source).read_bytes()
        names = pd.ExcelFile(io.BytesIO(raw)).sheet_names
        if "Header" in names and "Items" in names:
            return workbook_to_flat(raw)
        return read_table(raw, filename)
    return read_table(source, filename)


def template_xlsx_bytes() -> bytes:
    """4-sheet workbook template (same layout as Subcontracting_PO_Upload.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    s = SAMPLE_ROWS[0]
    sample = {
        "Header": {c: s.get(c, "") for c in SHEETS["Header"]},
        "Items": {c: s.get(c, "") for c in SHEETS["Items"]},
        "ScheduleLines": {"PO_Ref": "PO1", "PurchaseOrderItem": "10", "ScheduleLine": "1",
                          "ScheduleLineDeliveryDate": "2026-10-20",
                          "ScheduleLineOrderQuantity": "100"},
        "Components": {"PO_Ref": "PO1", "PurchaseOrderItem": "10", "ScheduleLine": "1",
                       "ReservationItem": "1", "Material": "RM019", "RequirementDate": "2026-10-20",
                       "QuantityInEntryUnit": "100", "EntryUnit": "KG", "Plant": "1810"},
    }
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="5B21B6")
    wb = Workbook()
    info = wb.active
    info.title = "Instructions"
    lines = [
        ("Subcontracting PO upload template (API_PURCHASEORDER_PROCESS_SRV, DEV)", True),
        ("Fill Header, Items, ScheduleLines and Components. Row 2 of each sheet is an example - "
         "replace it with your own data.", False),
        ("Header: one row per PO. PO_Ref is your own key.", False),
        ("Items: linked to Header by PO_Ref. Material is the finished good. Item category is "
         "always sent as 3 (subcontracting).", False),
        ("ScheduleLines: linked to Items by PO_Ref + PurchaseOrderItem.", False),
        ("Components: linked to schedule lines by PO_Ref + PurchaseOrderItem + ScheduleLine. "
         "One row per component material.", False),
        ("Dates: YYYY-MM-DD or DD.MM.YYYY. Item 10 and schedule line 1 are padded to 00010 / "
         "0001 automatically. Cells are text-formatted so leading zeros are kept.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=text)
        c.font = Font(name="Arial", bold=bold, size=13 if bold else 10)
        c.alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 110
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        for j, col in enumerate(cols, 1):
            h = ws.cell(row=1, column=j, value=col)
            h.font, h.fill = head_font, head_fill
            ws.column_dimensions[h.column_letter].width = max(12, len(col) + 2)
            v = ws.cell(row=2, column=j, value=sample[name].get(col, ""))
            v.font = Font(name="Arial")
            for r in range(2, 301):
                ws.cell(row=r, column=j).number_format = "@"
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Payloads
# ----------------------------------------------------------------------------
def build_documents(df) -> list[dict]:
    md = load_metadata(SERVICE)
    hb, ib = EntityBuilder(md, HEADER_TYPE), EntityBuilder(md, ITEM_TYPE)
    sb, cb = EntityBuilder(md, SCHED_TYPE), EntityBuilder(md, COMP_TYPE)
    pos: dict[str, dict] = {}

    for row in rows_of(df):
        errs: list[str] = []
        ref = row.get("PO_Ref") or f"LINE{row['_line']}"
        if ref not in pos:
            for f in ("CompanyCode", "PurchaseOrderType", "Supplier", "PurchasingOrganization",
                      "PurchasingGroup", "DocumentCurrency"):
                if not row.get(f):
                    errs.append(f"{f} is required on the first row of {ref}")
            header = hb.build(row, HEADER_FIELDS, errs)
            header["to_PurchaseOrderItem"] = {"results": []}
            pos[ref] = {"ref": ref, "lines": [], "errors": [], "payload": header, "items": {}}
        po = pos[ref]

        item_no = pad("PurchaseOrderItem", row.get("PurchaseOrderItem")
                      or str(10 * (len(po["items"]) + 1)))
        if item_no not in po["items"]:
            cat = row.get("PurchaseOrderItemCategory", "").upper()
            if cat not in ACCEPTED_CATEGORY_INPUTS:
                errs.append(f"item {item_no}: category {cat} is not subcontracting - use the "
                            f"Standard PO upload")
            for f in ("Material", "Plant", "OrderQuantity", "PurchaseOrderQuantityUnit",
                      "DeliveryDate"):
                if not row.get(f):
                    errs.append(f"{f} is required for item {item_no}")
            item = ib.build(dict(row, PurchaseOrderItem=item_no,
                                 PurchaseOrderItemCategory=SUBCONTRACTING_ITEM_CATEGORY),
                            ITEM_FIELDS, errs)
            item["to_ScheduleLine"] = {"results": []}
            po["payload"]["to_PurchaseOrderItem"]["results"].append(item)
            po["items"][item_no] = {"item": item, "scheds": {}, "row": row, "n_comp": 0}
        it = po["items"][item_no]

        sched_no = pad("ScheduleLine", row.get("ScheduleLine") or "1")
        if sched_no not in it["scheds"]:
            sl = sb.build(dict(row, ScheduleLine=sched_no), SCHED_MAP, errs,
                          defaults={"ScheduleLineOrderQuantity": row.get("OrderQuantity")
                                    or it["row"].get("OrderQuantity", "")})
            it["item"]["to_ScheduleLine"]["results"].append(sl)
            it["scheds"][sched_no] = sl
        sl = it["scheds"][sched_no]

        if row.get("ComponentMaterial"):
            for f in ("ComponentQuantity", "ComponentUnit"):
                if not row.get(f):
                    errs.append(f"{f} is required for component {row['ComponentMaterial']}")
            comps = sl.setdefault("to_SubcontractingComponent", {"results": []})["results"]
            comp = cb.build(row, COMP_MAP, errs, defaults={
                "ReservationItem": str(len(comps) + 1),
                "RequirementDate": row.get("DeliveryDate") or it["row"].get("DeliveryDate", ""),
                "Plant": row.get("Plant") or it["row"].get("Plant", "")})
            comps.append(comp)
            it["n_comp"] += 1

        po["lines"].append(row["_line"])
        po["errors"] += [f"line {row['_line']}: {e}" for e in errs]

    docs = []
    for po in pos.values():
        errors = list(po["errors"])
        if not ALLOW_BOM_EXPLOSION:
            for no, it in po["items"].items():
                if it["n_comp"] == 0:
                    errors.append(f"item {no}: no components - add ComponentMaterial rows "
                                  f"(or set ALLOW_BOM_EXPLOSION = True to let SAP explode the BOM)")
        errors += validate_payload(md, HEADER_TYPE, po["payload"])
        docs.append({"ref": po["ref"], "lines": ",".join(map(str, po["lines"])),
                     "errors": errors, "payload": po["payload"]})
    return docs


if __name__ == "__main__":
    run_cli(sys.modules[__name__])
