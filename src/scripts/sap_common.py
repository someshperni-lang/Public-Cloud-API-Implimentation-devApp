"""
Shared helpers for the SAP DEV upload scripts and dashboard.

* Metadata   - reads the EDMX-as-JSON files in ../metadata and exposes entity
               types, fields, max lengths and navigation properties.
* Conversion - turns spreadsheet text into SAP OData V2 values, driven by the
               field type in the metadata (dates -> /Date(ms)/, booleans,
               decimals as strings, leading zeros restored).
* Validation - checks every payload against the metadata before posting
               (unknown field, value too long, wrong type).
* Files      - reads CSV / XLSX as plain text so no numpy / pandas types ever
               reach a payload (the cause of "not JSON serializable").
* SAPClient  - CSRF handling, POST with one automatic retry, SAP error text.
"""
from __future__ import annotations

import io
import json
import math
import os
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
import requests

BASE_DIR = Path(__file__).resolve().parent.parent
METADATA_DIR = BASE_DIR / "metadata"

# ----------------------------------------------------------------------------
# Metadata
# ----------------------------------------------------------------------------
def _strip(k: str) -> str:
    return k.split("}")[-1]


def _norm(o):
    if isinstance(o, dict):
        out = {}
        for k, v in o.items():
            if k == "@attributes":
                out["@"] = {_strip(a): b for a, b in v.items()}
            else:
                out[_strip(k)] = _norm(v)
        return out
    if isinstance(o, list):
        return [_norm(x) for x in o]
    return o


def _aslist(x):
    if x is None:
        return []
    return x if isinstance(x, list) else [x]


class ServiceMetadata:
    """One OData service, parsed from an EDMX file exported as JSON."""

    def __init__(self, path: Path):
        self.path = Path(path)
        raw = _norm(json.loads(self.path.read_text(encoding="utf-8")))
        schema = _aslist(raw["Edmx"]["DataServices"]["Schema"])[0]
        self.service = schema["@"]["Namespace"]
        self.entity_types: dict[str, dict] = {}
        for et in _aslist(schema.get("EntityType")):
            a = et["@"]
            keys = [r["@"]["Name"] for r in _aslist((et.get("Key") or {}).get("PropertyRef"))]
            props = {p["@"]["Name"]: p["@"] for p in _aslist(et.get("Property"))}
            navs = {n["@"]["Name"]: n["@"] for n in _aslist(et.get("NavigationProperty"))}
            self.entity_types[a["Name"]] = {"label": a.get("label", ""), "keys": keys,
                                            "props": props, "navs": navs}
        self.assoc_ends: dict[tuple, dict] = {}
        for asn in _aslist(schema.get("Association")):
            name = f'{self.service}.{asn["@"]["Name"]}'
            for end in _aslist(asn.get("End")):
                self.assoc_ends[(name, end["@"]["Role"])] = end["@"]
        container = _aslist(schema.get("EntityContainer"))[0]
        self.entity_sets: dict[str, dict] = {}
        for es in _aslist(container.get("EntitySet")):
            a = es["@"]
            a["EntityTypeName"] = a["EntityType"].split(".")[-1]
            self.entity_sets[a["Name"]] = a
        self.function_imports = [f["@"]["Name"] for f in _aslist(container.get("FunctionImport"))]

    def type_of_set(self, entity_set: str) -> str:
        return self.entity_sets[entity_set]["EntityTypeName"]

    def nav_target(self, entity_type: str, nav: str) -> tuple[str, bool]:
        """Return (target entity type, is_collection) for a navigation property."""
        n = self.entity_types[entity_type]["navs"][nav]
        end = self.assoc_ends[(n["Relationship"], n["ToRole"])]
        return end["Type"].split(".")[-1], end.get("Multiplicity") == "*"

    def field(self, entity_type: str, name: str) -> dict | None:
        return self.entity_types[entity_type]["props"].get(name)

    def fields_table(self, entity_type: str) -> list[dict]:
        et = self.entity_types[entity_type]
        rows = []
        for name, p in et["props"].items():
            rows.append({
                "Field": name,
                "Label": p.get("label", ""),
                "Type": p.get("Type", "").replace("Edm.", ""),
                "MaxLength": p.get("MaxLength", ""),
                "Key": "X" if name in et["keys"] else "",
                "Creatable": "" if p.get("creatable") == "false" else "X",
                "Updatable": "" if p.get("updatable") == "false" else "X",
                "Description": p.get("quickinfo", ""),
            })
        return rows


_METADATA_CACHE: dict[str, ServiceMetadata] = {}


def load_metadata(service: str, metadata_dir: Path | None = None) -> ServiceMetadata | None:
    """Find and parse the metadata JSON for a service (e.g. API_PRODUCT_SRV)."""
    if service in _METADATA_CACHE:
        return _METADATA_CACHE[service]
    folder = Path(metadata_dir or METADATA_DIR)
    if not folder.exists():
        return None
    for f in sorted(folder.glob("*.json")):
        if re.search(rf"(^|_){re.escape(service)}(_\d+)?\.json$", f.name):
            md = ServiceMetadata(f)
            _METADATA_CACHE[service] = md
            return md
    return None


def all_metadata(metadata_dir: Path | None = None) -> dict[str, ServiceMetadata]:
    folder = Path(metadata_dir or METADATA_DIR)
    out = {}
    for f in sorted(folder.glob("*.json")):
        try:
            md = ServiceMetadata(f)
            out[md.service] = md
        except Exception:
            continue
    return out


# ----------------------------------------------------------------------------
# Value conversion
# ----------------------------------------------------------------------------
# Fields where Excel tends to strip leading zeros, with the length to pad to.
# Only purely numeric values are padded.
PAD_FIELDS = {
    "PurchaseOrderItem": 5, "PurchasingDocumentItem": 5, "PurchaseRequisitionItem": 5,
    "PurchaseContractItem": 5, "ScheduleLine": 4, "ReservationItem": 4,
    "AccountAssignmentNumber": 2, "MaterialDocumentItem": 4, "ReferenceDocumentItem": 4,
    "ManufacturingOrderItem": 4, "DeliveryItem": 6, "SalesOrderItem": 6,
    "GoodsMovementCode": 2, "PurchasingGroup": 3, "MRPResponsible": 3,
    "CostCenter": 10, "ProfitCenter": 10, "GLAccount": 10, "Division": 2,
    "ProductDistributionChnl": 2, "ManufacturingOrder": 12, "OrderID": 12,
}


def is_blank(v) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    s = str(v).strip()
    return s == "" or s.lower() in ("nan", "none", "nat", "null")


def clean(v) -> str:
    """Return a trimmed string, '' for blanks, and drop Excel's trailing '.0'."""
    if is_blank(v):
        return ""
    s = str(v).strip()
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".")[0]
    return s


def pad(field: str, value: str) -> str:
    n = PAD_FIELDS.get(field)
    if n and value.isdigit() and len(value) < n:
        return value.zfill(n)
    return value


_DATE_FORMATS = ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y",
                 "%Y/%m/%d", "%Y%m%d", "%d-%b-%Y", "%d %b %Y", "%Y-%m-%dT%H:%M:%S"]


def parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = clean(v)
    m = re.fullmatch(r"/Date\((-?\d+)\)/", s)
    if m:
        return datetime.fromtimestamp(int(m.group(1)) / 1000, tz=timezone.utc).date()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    if re.fullmatch(r"\d{5}(\.\d+)?", s):          # Excel serial number
        return (datetime(1899, 12, 30) + timedelta(days=float(s))).date()
    raise ValueError(f"unrecognised date '{s}' (use YYYY-MM-DD or DD.MM.YYYY)")


def sap_date(v) -> str:
    d = parse_date(v)
    ms = int(datetime(d.year, d.month, d.day, tzinfo=timezone.utc).timestamp() * 1000)
    return f"/Date({ms})/"


def sap_time(v) -> str:
    """Edm.Time -> PT10H30M00S"""
    s = clean(v)
    if s.startswith("PT"):
        return s
    m = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
    if not m:
        raise ValueError(f"unrecognised time '{s}' (use HH:MM or HH:MM:SS)")
    return f"PT{int(m.group(1)):02d}H{int(m.group(2)):02d}M{int(m.group(3) or 0):02d}S"


def sap_bool(v) -> bool:
    s = clean(v).lower()
    if s in ("true", "x", "yes", "y", "1"):
        return True
    if s in ("false", "no", "n", "0", "-"):
        return False
    raise ValueError(f"unrecognised yes/no value '{v}' (use TRUE/FALSE or X)")


def sap_decimal(v) -> str:
    s = clean(v).replace(" ", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")              # 1,5 -> 1.5
    else:
        s = s.replace(",", "")               # 1,250.00 -> 1250.00
    try:
        Decimal(s)
    except InvalidOperation:
        raise ValueError(f"'{v}' is not a number")
    if "e" in s.lower():
        s = format(Decimal(s), "f")
    return s


def convert(edm_type: str, field: str, value):
    """Convert one spreadsheet value to the JSON value SAP expects for edm_type."""
    t = (edm_type or "Edm.String").replace("Edm.", "")
    if t == "DateTime" or t == "DateTimeOffset":
        return sap_date(value)
    if t == "Time":
        return sap_time(value)
    if t == "Boolean":
        return sap_bool(value)
    if t in ("Decimal", "Double", "Single"):
        return sap_decimal(value)
    if t in ("Int16", "Int32", "Int64", "Byte", "SByte"):
        return int(Decimal(sap_decimal(value)))
    s = clean(value)
    return pad(field, s)


class EntityBuilder:
    """Builds the dict for one entity from a row, converting by metadata type."""

    def __init__(self, md: ServiceMetadata | None, entity_type: str):
        self.md = md
        self.entity_type = entity_type

    def edm_type(self, field: str) -> str:
        if self.md:
            p = self.md.field(self.entity_type, field)
            if p:
                return p.get("Type", "Edm.String")
        return "Edm.String"

    def build(self, row: dict, mapping: dict[str, str] | list[str], errors: list[str],
              defaults: dict | None = None) -> dict:
        """mapping: {sap_field: column_name} or [field,...] where column == field."""
        if isinstance(mapping, list):
            mapping = {f: f for f in mapping}
        out = {}
        for sap_field, col in mapping.items():
            raw = row.get(col, "")
            if is_blank(raw) and defaults and sap_field in defaults:
                raw = defaults[sap_field]
            if is_blank(raw):
                continue
            try:
                out[sap_field] = convert(self.edm_type(sap_field), sap_field, raw)
            except ValueError as e:
                errors.append(f"{col}: {e}")
        return out


# ----------------------------------------------------------------------------
# Payload validation against metadata
# ----------------------------------------------------------------------------
def validate_payload(md: ServiceMetadata | None, entity_type: str, payload: dict,
                     path: str = "") -> list[str]:
    """Check field names, max lengths and types, following deep-insert navigations."""
    if md is None:
        return []
    errs = []
    et = md.entity_types.get(entity_type)
    if not et:
        return [f"{path or entity_type}: entity type not in metadata"]
    for k, v in payload.items():
        loc = f"{path}{k}"
        if k in et["navs"]:
            target, _ = md.nav_target(entity_type, k)
            children = v.get("results", []) if isinstance(v, dict) and "results" in v else (
                v if isinstance(v, list) else [v])
            for i, child in enumerate(children):
                errs += validate_payload(md, target, child, f"{loc}[{i}].")
            continue
        p = et["props"].get(k)
        if p is None:
            errs.append(f"{loc}: field does not exist in {entity_type}")
            continue
        if p.get("creatable") == "false" and k not in et["keys"]:
            errs.append(f"{loc}: field is read-only (not creatable)")
        t = p.get("Type", "")
        if t == "Edm.String" and p.get("MaxLength") and isinstance(v, str):
            if len(v) > int(p["MaxLength"]):
                errs.append(f"{loc}: '{v}' is longer than {p['MaxLength']} characters")
        if t == "Edm.Boolean" and not isinstance(v, bool):
            errs.append(f"{loc}: must be true/false")
        if t == "Edm.DateTime" and not (isinstance(v, str) and v.startswith("/Date(")):
            errs.append(f"{loc}: must be an SAP date /Date(ms)/")
        if t == "Edm.Decimal" and p.get("Precision") and isinstance(v, str):
            digits = v.lstrip("-").replace(".", "")
            scale = int(p.get("Scale") or 0)
            if "." in v and len(v.split(".")[1]) > scale:
                errs.append(f"{loc}: '{v}' has more than {scale} decimal places")
            elif len(digits.lstrip("0")) > int(p["Precision"]):
                errs.append(f"{loc}: '{v}' exceeds {p['Precision']} digits")
    return errs


# ----------------------------------------------------------------------------
# JSON safety (belt and braces)
# ----------------------------------------------------------------------------
def make_json_safe(obj):
    if isinstance(obj, dict):
        return {k: make_json_safe(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [make_json_safe(v) for v in obj]
    if obj is None or obj is pd.NaT:
        return None
    try:
        import numpy as np
        if isinstance(obj, np.bool_):
            return bool(obj)
        if isinstance(obj, np.integer):
            return int(obj)
        if isinstance(obj, np.floating):
            return None if math.isnan(obj) else float(obj)
    except ImportError:
        pass
    if isinstance(obj, float) and math.isnan(obj):
        return None
    if isinstance(obj, Decimal):
        return format(obj, "f")
    if isinstance(obj, (datetime, date)):
        return sap_date(obj)
    if isinstance(obj, (str, int, float, bool)):
        return obj
    return str(obj)


# ----------------------------------------------------------------------------
# File reading
# ----------------------------------------------------------------------------
def read_table(source, filename: str | None = None, sheet: str | None = None) -> pd.DataFrame:
    """Read CSV or XLSX (path, bytes or file-like) with every cell as plain text."""
    name = (filename or getattr(source, "name", None) or str(source)).lower()
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        xl = pd.ExcelFile(source)
        use = sheet if sheet in xl.sheet_names else (
            "Upload" if "Upload" in xl.sheet_names else xl.sheet_names[0])
        df = pd.read_excel(xl, sheet_name=use, dtype=str, keep_default_na=False)
    else:
        raw = source.read() if hasattr(source, "read") else Path(source).read_bytes()
        text = raw.decode("utf-8-sig", errors="replace") if isinstance(raw, bytes) else raw
        sep = ";" if text.splitlines()[0].count(";") > text.splitlines()[0].count(",") else ","
        df = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False, sep=sep)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[:, [c for c in df.columns if c and not c.startswith("Unnamed")]]
    df = df.map(clean) if hasattr(df, "map") else df.applymap(clean)
    df = df[(df != "").any(axis=1)].reset_index(drop=True)
    return df


def rows_of(df: pd.DataFrame) -> list[dict]:
    """Rows as plain dicts, with the spreadsheet line number in '_line'."""
    rows = []
    for i, r in enumerate(df.to_dict(orient="records")):
        r = {k: clean(v) for k, v in r.items()}
        r["_line"] = i + 2
        rows.append(r)
    return rows


def template_xlsx(columns: list[str], sample_rows: list[dict], notes: dict[str, str],
                  title: str) -> bytes:
    """Excel template: 'Upload' sheet (text-formatted) + 'Instructions' sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Upload"
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    body_font = Font(name="Arial")
    for c, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=col)
        cell.font, cell.fill = head_font, head_fill
        ws.column_dimensions[cell.column_letter].width = max(12, len(col) + 2)
    for r, row in enumerate(sample_rows, 2):
        for c, col in enumerate(columns, 1):
            cell = ws.cell(row=r, column=c, value=row.get(col, ""))
            cell.font = body_font
    for c in range(1, len(columns) + 1):          # text format keeps leading zeros
        for r in range(2, 501):
            ws.cell(row=r, column=c).number_format = "@"
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    info["A1"] = title
    info["A1"].font = Font(name="Arial", bold=True, size=13)
    info["A2"] = ("Fill the 'Upload' sheet. Rows 2+ hold example values - replace them with "
                  "your own. Cells are text-formatted so leading zeros are kept. "
                  "Dates: YYYY-MM-DD or DD.MM.YYYY. Yes/no: TRUE/FALSE or X.")
    info["A2"].alignment = Alignment(wrap_text=True)
    info.merge_cells("A2:C2")
    info.row_dimensions[2].height = 45
    info["A4"], info["B4"] = "Column", "What to enter"
    for c in ("A4", "B4"):
        info[c].font = head_font
        info[c].fill = head_fill
    for i, col in enumerate(columns, 5):
        info.cell(row=i, column=1, value=col).font = Font(name="Arial", bold=True)
        info.cell(row=i, column=2, value=notes.get(col, "")).font = body_font
    info.column_dimensions["A"].width = 32
    info.column_dimensions["B"].width = 95
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_csv(columns: list[str], sample_rows: list[dict]) -> bytes:
    df = pd.DataFrame([{c: r.get(c, "") for c in columns} for r in sample_rows], columns=columns)
    return df.to_csv(index=False).encode("utf-8-sig")


# ----------------------------------------------------------------------------
# SAP client
# ----------------------------------------------------------------------------
def load_env_file(path: Path | None = None):
    """Minimal .env loader (no extra dependency)."""
    p = Path(path or BASE_DIR / ".env")
    if not p.exists():
        return
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


# DEV system (SAP S/4HANA Cloud). Each script's BASE_URL is built from this.
DEV_HOST = "https://my440263-api.s4hana.cloud.sap"


def service_url(service: str, host: str = DEV_HOST) -> str:
    return f"{host.rstrip('/')}/sap/opu/odata/sap/{service}"


def credentials_from_env() -> tuple[str, str]:
    """Communication user / password from .env or the environment."""
    load_env_file()
    return (os.getenv("SAP_API_USER", ""), os.getenv("SAP_API_PASSWORD", ""))


def extract_error(resp) -> str:
    try:
        err = resp.json()["error"]
        msg = err["message"]["value"]
        details = [d.get("message", "") for d in err.get("innererror", {}).get("errordetails", [])]
        details = [d for d in dict.fromkeys(details) if d and d != msg]
        return " | ".join([msg] + details)
    except Exception:
        text = getattr(resp, "text", "") or ""
        m = re.search(r"<message[^>]*>(.*?)</message>", text, re.S)
        return (m.group(1) if m else text)[:500] or f"HTTP {resp.status_code} with empty body"


class SAPClient:
    """Session against one OData service root, e.g.
    https://my440263-api.s4hana.cloud.sap/sap/opu/odata/sap/API_PURCHASEORDER_PROCESS_SRV"""

    def __init__(self, base_url: str, user: str, password: str, timeout: int = 120):
        self.base = base_url.strip().rstrip("/")
        self.timeout = timeout
        self.session = requests.Session()
        self.session.auth = (user, password)
        self.session.headers.update({"Accept": "application/json",
                                     "Content-Type": "application/json"})
        self.csrf = None

    def fetch_csrf(self):
        r = self.session.get(f"{self.base}/", headers={"X-CSRF-Token": "Fetch"},
                             timeout=self.timeout)
        if r.status_code == 401:
            raise RuntimeError("401 Unauthorized - check the communication user and password")
        if r.status_code == 403:
            raise RuntimeError("403 Forbidden - the communication arrangement does not "
                               "allow this user to call this service")
        if r.status_code == 404:
            raise RuntimeError(f"404 Not found - check the service URL {self.base}")
        r.raise_for_status()
        self.csrf = r.headers.get("x-csrf-token")
        if not self.csrf:
            raise RuntimeError("SAP did not return a CSRF token")

    def test(self) -> tuple[bool, str]:
        try:
            self.fetch_csrf()
            return True, "Connected"
        except requests.exceptions.ConnectionError:
            return False, f"Cannot reach {self.base}"
        except Exception as e:
            return False, str(e)

    def post(self, entity_set: str, payload: dict):
        if not self.csrf:
            self.fetch_csrf()
        resp = self._post(entity_set, payload)
        if resp.status_code == 403 and "required" in resp.headers.get("x-csrf-token", "").lower():
            self.fetch_csrf()                      # token expired - fetch once and retry
            resp = self._post(entity_set, payload)
        return resp

    def get(self, path: str):
        return self.session.get(f"{self.base}/{path}", timeout=self.timeout)

    def patch(self, path: str, payload: dict, etag: str | None = None):
        """PATCH with CSRF (and If-Match), retrying once if the token expired."""
        if not self.csrf:
            self.fetch_csrf()

        def do():
            return self.session.patch(f"{self.base}/{path}",
                                      headers={"X-CSRF-Token": self.csrf, "If-Match": etag or "*"},
                                      data=json.dumps(make_json_safe(payload)),
                                      timeout=self.timeout)
        resp = do()
        if resp.status_code == 403 and "required" in resp.headers.get("x-csrf-token", "").lower():
            self.fetch_csrf()
            resp = do()
        return resp

    def _post(self, entity_set, payload):
        return self.session.post(f"{self.base}/{entity_set}",
                                 headers={"X-CSRF-Token": self.csrf},
                                 data=json.dumps(make_json_safe(payload)), timeout=self.timeout)


def module_template_xlsx(module) -> bytes:
    """A script may ship its own Excel template (e.g. the 4-sheet subcontracting one)."""
    if hasattr(module, "template_xlsx_bytes"):
        return module.template_xlsx_bytes()
    return template_xlsx(module.COLUMNS, module.SAMPLE_ROWS, module.COLUMN_NOTES, module.TITLE)


def module_read_upload(module, source, filename: str) -> pd.DataFrame:
    """A script may read its own upload layout (e.g. a multi-sheet workbook)."""
    if hasattr(module, "read_upload"):
        return module.read_upload(source, filename)
    return read_table(source, filename)


RESULT_COLUMNS = ["Timestamp", "Module", "Group", "Items", "Status", "HTTP",
                  "Document", "Year", "Detail"]


def post_documents(client: SAPClient, module, documents: list[dict], stop_on_fail=False,
                   progress=None, pause: float = 0.3) -> list[dict]:
    """POST every valid document. Rows use RESULT_COLUMNS (same as the main dashboard)."""
    import time

    rows = []
    total = len(documents)
    for i, doc in enumerate(documents, 1):
        row = dict.fromkeys(RESULT_COLUMNS, "")
        row.update(Timestamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                   Module=module.TITLE, Group=doc["ref"], Items=module.count_items(doc["payload"]))
        if doc["errors"]:
            row.update(Status="Invalid", Detail=" | ".join(doc["errors"]))
        else:
            try:
                note = ""
                if hasattr(module, "post_document"):          # script-specific write (e.g. upsert)
                    r, note = module.post_document(client, doc["payload"])
                else:
                    r = client.post(module.ENTITY_SET, doc["payload"])
                row["HTTP"] = str(r.status_code)
                if 200 <= r.status_code < 300:
                    try:
                        d = r.json().get("d", {}) if r.text.strip() else {}
                    except ValueError:
                        d = {}
                    doc_no = str(d.get(module.DOC_FIELD, "") or "")
                    if not doc_no and getattr(module, "GROUP_AS_DOC", False):
                        doc_no = doc["ref"]
                    row.update(Status="Success", Document=doc_no, Detail=note,
                               Year=str(d.get(module.YEAR_FIELD, "")) if module.YEAR_FIELD else "")
                else:
                    row.update(Status=f"Failed ({r.status_code})", Detail=extract_error(r))
            except Exception as e:
                row.update(Status="Error", Detail=f"{type(e).__name__}: {e}"[:1000])
            time.sleep(pause)
        rows.append(row)
        if progress:
            progress(i, total, row)
        if stop_on_fail and row["Status"] not in ("Success", "Invalid"):
            break
    return rows


# ----------------------------------------------------------------------------
# Common CLI
# ----------------------------------------------------------------------------
def run_cli(module, argv=None):
    """Shared command line for every upload script."""
    import argparse
    import sys

    ap = argparse.ArgumentParser(description=module.TITLE)
    ap.add_argument("-f", "--file", help="CSV or XLSX upload file")
    ap.add_argument("--dry-run", action="store_true", help="Build + validate payloads only")
    ap.add_argument("--template", choices=["csv", "xlsx"], help="Write a blank template and exit")
    ap.add_argument("--base-url", default=module.BASE_URL,
                    help=f"Service root (default: {module.BASE_URL})")
    ap.add_argument("--stop-on-fail", action="store_true", help="Stop at the first failure")
    ap.add_argument("--out", help="Results CSV path (default: results/<slug>_<time>.csv)")
    args = ap.parse_args(argv)

    if args.template:
        name = f"{module.SLUG}_template.{args.template}"
        data = (module_template_xlsx(module) if args.template == "xlsx"
                else template_csv(module.COLUMNS, module.SAMPLE_ROWS))
        Path(name).write_bytes(data)
        print(f"Template written: {name}")
        return
    if not args.file:
        ap.error("--file is required (or use --template)")

    docs = module.build_documents(module_read_upload(module, args.file, args.file))
    out_dir = Path("payloads") / module.SLUG
    out_dir.mkdir(parents=True, exist_ok=True)
    bad = 0
    for d in docs:
        (out_dir / f"{re.sub(r'[^A-Za-z0-9_-]', '_', d['ref'])}.json").write_text(
            json.dumps(d["payload"], indent=2), encoding="utf-8")
        bad += bool(d["errors"])
        print(f"--- {d['ref']} (lines {d['lines']}): {'INVALID' if d['errors'] else 'ok'}")
        for e in d["errors"]:
            print(f"    {e}")
    print(f"Built {len(docs)} document(s), {bad} invalid. Payloads in {out_dir}/")
    if args.dry_run:
        return

    user, password = credentials_from_env()
    if not (user and password):
        sys.exit("Set SAP_API_USER and SAP_API_PASSWORD in .env or the environment.")
    client = SAPClient(args.base_url, user, password)
    rows = post_documents(client, module, docs, args.stop_on_fail,
                          progress=lambda i, n, r: print(f"[{i}/{n}] {r['Group']}: {r['Status']} "
                                                         f"{r['Document']} {r['Detail'][:200]}"))
    out = Path(args.out) if args.out else (
        BASE_DIR / "results" / f"{module.SLUG}_{datetime.now():%Y%m%d_%H%M%S}.csv")
    out.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows, columns=RESULT_COLUMNS).to_csv(out, index=False)
    print(f"Results written to {out}")
